"""Resume scorer backend. Run with:  uvicorn app:app --reload"""
import io
import math
import re
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Optional

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader

BASE = Path(__file__).parent
MAX_BYTES = 5 * 1024 * 1024

# ---------------------------------------------------------------- skills ----
# First name is shown in the UI, the rest are aliases. Add your own lines.
SKILL_LINES = """
Python
Java
JavaScript|JS|ECMAScript
TypeScript
C++|CPP
C#|CSharp
Golang
Rust
PHP
Ruby
Kotlin
Swift
Scala
SQL
MATLAB
Bash|Shell Scripting|Shell Script
HTML|HTML5
CSS|CSS3
SASS|SCSS
React|React.js|ReactJS
React Native
Angular|AngularJS
Vue.js|Vue|VueJS
Next.js|NextJS
Node.js|NodeJS|Node JS
Express.js|ExpressJS|Express JS
EJS
Django
Flask
FastAPI
Spring Boot|SpringBoot
Hibernate
.NET|dotnet|ASP.NET
jQuery
Bootstrap
Tailwind CSS|Tailwind|TailwindCSS
Redux
Webpack
REST API|REST APIs|RESTful|RESTful API|RESTful APIs
GraphQL
WebSockets|WebSocket
JWT
OAuth
Postman
Swagger|OpenAPI
MVC
Pandas
NumPy
SciPy
scikit-learn|sklearn|scikit learn
TensorFlow
PyTorch
Keras
OpenCV
NLP|Natural Language Processing
Machine Learning|ML
Deep Learning
Computer Vision
Artificial Intelligence|AI
Generative AI|GenAI|Gen AI
LLMs|LLM|Large Language Models
LangChain
Data Analysis|Data Analytics
Data Visualization|Data Visualisation
Tableau
Power BI|PowerBI
Excel|Microsoft Excel|MS Excel
Statistics|Statistical Analysis
ETL
Spark|PySpark|Apache Spark
Hadoop
Kafka|Apache Kafka
Airflow|Apache Airflow
RabbitMQ
Celery
MongoDB|Mongo DB
Mongoose
MySQL
PostgreSQL|Postgres
SQLite
Redis
Oracle
DynamoDB
Elasticsearch
Firebase
NoSQL
AWS|Amazon Web Services
Azure|Microsoft Azure
GCP|Google Cloud|Google Cloud Platform
Docker
Kubernetes|K8s
Terraform
Jenkins
CI/CD|CICD|CI CD|Continuous Integration
GitHub Actions
Git
GitHub
GitLab
Linux
Nginx
Ansible
Maven
Pytest
JUnit
Selenium
Jest
Unit Testing
Android
iOS
Flutter
Agile
Scrum
Jira
Project Management
Microservices|Microservice
OOP|Object Oriented Programming|Object-Oriented Programming|Object Oriented
Data Structures
Algorithms
System Design
Figma
UI/UX|UX Design|UI Design
SEO
Salesforce
SAP
Power Automate
Web Scraping
"""

SKILLS = {}
for _line in SKILL_LINES.strip().splitlines():
    _aliases = [a.strip() for a in _line.split("|") if a.strip()]
    if _aliases:
        SKILLS[_aliases[0]] = _aliases


def _alias_regex(alias):
    return r"[\s\-]+".join(re.escape(p) for p in alias.split())


SKILL_PATTERNS = {
    name: re.compile(
        r"(?<![A-Za-z0-9+#.])(?:%s)(?![A-Za-z+#])"
        % "|".join(sorted((_alias_regex(a) for a in aliases), key=len, reverse=True)),
        re.I,
    )
    for name, aliases in SKILLS.items()
}


def find_skills(text):
    found = {}
    for name, rx in SKILL_PATTERNS.items():
        forms = {m.group(0).lower() for m in rx.finditer(text)}
        if forms:
            found[name] = forms
    return found


# ---------------------------------------------------------- file reading ----
class ParseError(ValueError):
    """Message is safe to show to the user."""


def normalise(text):
    text = text.replace("\x00", " ").replace("\u00a0", " ")
    for bullet in "•●▪◦■□➢➤►·":
        text = text.replace(bullet, "- ")
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _pdf_text(data):
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise ParseError("This PDF is password-protected. Remove the password and upload it again.")
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"Could not read this PDF ({type(exc).__name__}).") from exc


def _docx_text(data):
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("Could not read this DOCX file. Old .doc files are not supported.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:
                    cells.append(t)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename, data):
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        raw = _pdf_text(data)
    elif ext == ".docx":
        raw = _docx_text(data)
    elif ext in (".txt", ".md"):
        raw = data.decode("utf-8", errors="ignore")
    else:
        raise ParseError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    text = normalise(raw)
    if len(text) < 30:
        raise ParseError("No readable text found. If this is a scanned PDF, upload a text-based version.")
    return text


# ------------------------------------------------------ resume structure ----
SECTION_HEADINGS = {
    "experience": ["experience", "work experience", "professional experience", "employment",
                   "employment history", "work history", "internships", "internship experience",
                   "career history"],
    "education": ["education", "academic background", "academics", "qualifications",
                  "educational qualifications", "academic qualifications"],
    "skills": ["skills", "technical skills", "key skills", "core competencies", "skills summary",
               "technologies", "tech stack", "technical proficiencies"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "summary": ["summary", "professional summary", "profile", "objective", "career objective", "about me"],
    "certifications": ["certifications", "certificates", "courses", "licenses"],
    "achievements": ["achievements", "awards", "honors", "accomplishments", "extracurricular"],
}
_HEADING_LOOKUP = {a: sec for sec, al in SECTION_HEADINGS.items() for a in al}


def _heading_key(line):
    if not line or len(line) > 40:
        return None
    norm = " ".join(re.sub(r"[^a-z& ]", "", line.lower()).split())
    return _HEADING_LOOKUP.get(norm)


def split_sections(text):
    sections, current = {}, "header"
    for line in text.splitlines():
        key = _heading_key(line.strip())
        if key:
            current = key
            sections.setdefault(current, [])
        else:
            sections.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(
    r"(?<!\d)(?:\+\d{1,3}[\s\-]?)?(?:\d{5}[\s\-]?\d{5}|\(?\d{3}\)?[\s\-.]?\d{3}[\s\-.]?\d{4})(?!\d)"
)
LINK_RE = re.compile(r"(?:https?://)?(?:www\.)?(?:linkedin\.com/in/|github\.com/)[A-Za-z0-9_\-./]+", re.I)


def _guess_name(text):
    for line in text.splitlines()[:8]:
        s = line.strip()
        if not s or "@" in s or re.search(r"\d", s) or len(s) > 40:
            continue
        if s.lower() in {"resume", "curriculum vitae", "cv"} or _heading_key(s):
            continue
        words = s.replace(",", " ").split()
        if 2 <= len(words) <= 4 and all(re.fullmatch(r"[A-Za-z][A-Za-z.\-']*", w) for w in words):
            return s.title() if s.isupper() else s
    return None


def extract_contact(text):
    email, phone = EMAIL_RE.search(text), PHONE_RE.search(text)
    links = []
    for m in LINK_RE.finditer(text):
        link = m.group(0).rstrip("./")
        if link.lower() not in [x.lower() for x in links]:
            links.append(link)
    return {
        "name": _guess_name(text),
        "email": email.group(0) if email else None,
        "phone": phone.group(0).strip() if phone else None,
        "links": links[:4],
    }


_MONTHS = {m: i for i, m in enumerate("jan feb mar apr may jun jul aug sep oct nov dec".split(), 1)}
_MON = "jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec"
_RANGE = re.compile(
    r"(?:(?P<m1>%s)[a-z]*\.?,?\s+)?(?P<y1>(?:19|20)\d{2})\s*(?:-|–|—|to|till|until)\s*"
    r"(?:(?P<m2>%s)[a-z]*\.?,?\s+)?(?P<y2>(?:19|20)\d{2}|present|current|now|ongoing|today)" % (_MON, _MON),
    re.I,
)


def estimate_experience_years(text, sections):
    body = sections.get("experience") or "\n".join(
        v for k, v in sections.items() if k not in {"education", "certifications"}
    )
    today = date.today()
    spans = []
    for m in _RANGE.finditer(body):
        start = int(m.group("y1")) * 12 + _MONTHS.get((m.group("m1") or "").lower(), 6)
        end_raw = m.group("y2").lower()
        if end_raw.isdigit():
            end = int(end_raw) * 12 + _MONTHS.get((m.group("m2") or "").lower(), 6)
        else:
            end = today.year * 12 + today.month
        if 0 <= end - start <= 45 * 12:
            spans.append((start, end))
    spans.sort()
    total, cur_s, cur_e = 0, None, None
    for s, e in spans:
        if cur_e is None or s > cur_e:
            if cur_e is not None:
                total += cur_e - cur_s
            cur_s, cur_e = s, e
        else:
            cur_e = max(cur_e, e)
    if cur_e is not None:
        total += cur_e - cur_s
    explicit = [
        float(x)
        for x in re.findall(
            r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years?|yrs?)\s+(?:of\s+)?(?:[a-z\-]+\s+){0,3}?experience",
            text, re.I)
    ]
    return round(min(40.0, max([total / 12] + explicit)), 1)


_REQ_YEARS = [
    re.compile(r"(\d{1,2})\s*\+?\s*(?:(?:-|–|to)\s*\d{1,2}\s*\+?\s*)?(?:years?|yrs?)(?:\s+of)?(?:\s+[a-z\-/]+){0,4}?\s+(?:experience|exp)\b", re.I),
    re.compile(r"experience[^.\n]{0,60}?(\d{1,2})\s*\+?\s*(?:years?|yrs?)", re.I),
    re.compile(r"(?:minimum|at least|min\.?)\s*(?:of\s*)?(\d{1,2})\s*\+?\s*(?:years?|yrs?)", re.I),
]


def jd_required_years(text):
    vals = [int(m.group(1)) for rx in _REQ_YEARS for m in rx.finditer(text) if int(m.group(1)) <= 25]
    return max(vals) if vals else None


_DEGREES = [
    (4, "Doctorate", r"ph\.?\s?d|doctorate|doctoral"),
    (3, "Master's", r"master['’]?s?|m\.?\s?tech|m\.?\s?sc|mba|mca|m\.\s?s\.|post[\s\-]?graduate|pgdm"),
    (2, "Bachelor's", r"bachelor['’]?s?|b\.?\s?tech|b\.?\s?e\.|b\.?\s?sc|bca|b\.?\s?com|b\.\s?a\.|b\.\s?s\.|undergraduate"),
    (1, "Diploma", r"diploma"),
]
_DEGREE_RX = [
    (lvl, label, re.compile(r"(?<![a-z0-9])(?:%s)(?![a-z0-9])" % pat, re.I)) for lvl, label, pat in _DEGREES
]


def degree_levels(text):
    return [(lvl, label) for lvl, label, rx in _DEGREE_RX if rx.search(text)]


# ------------------------------------------------- keywords & similarity ----
TOKEN_RE = re.compile(r"[a-z][a-z0-9+#]*(?:[.\-][a-z0-9+#]+)*")
STOPWORDS = set("""a an the and or but if of to in on at by for with from as is are was were be been being this that
these those it its we you your our their they he she him her his hers i me my not no nor so than then too very can
will just do does did done have has had having would should could may might must shall about above across after again
all also am any because before below between both each few more most other over own same some such into through under
until up while who whom what which when where why how out off only once here there per via""".split())


def stem(tok):
    t = tok
    if len(t) > 4 and t.endswith("ies"):
        t = t[:-3] + "y"
    elif len(t) > 5 and t.endswith("ing"):
        t = t[:-3]
    elif len(t) > 4 and t.endswith("ed"):
        t = t[:-2]
    elif len(t) > 3 and t.endswith("s") and not t.endswith("ss"):
        t = t[:-1]
    if len(t) > 4 and t.endswith("e"):
        t = t[:-1]
    return t


GENERIC = {stem(w) for w in """experience experienced work working team teams ability able strong strongly skills skill
knowledge understanding required requirement requirements preferred preferably plus role roles responsibilities
responsibility responsible job company candidate candidates looking opportunity join including include includes etc
good great excellent years year yrs minimum least new using use used within related relevant position positions apply
application applicants ideal proven demonstrated hands-on develop developing development building build create creating
support supporting ensure ensuring provide providing help helping successful seeking well will""".split()}


def _tokens(text):
    return TOKEN_RE.findall(text.lower())


def _chunks(text):
    return re.split(r"[\n;:|,()]+|\.(?=\s|$)", text)


def _stem_sets(text):
    uni, bi = set(), set()
    for chunk in _chunks(text):
        prev = None
        for t in _tokens(chunk):
            if t in STOPWORDS:
                prev = None
                continue
            s = stem(t)
            uni.add(s)
            if prev is not None:
                bi.add((prev, s))
            prev = s
    return uni, bi


def extract_keywords(jd, skip, n_uni=25, n_bi=8):
    uni, bi, surface = Counter(), Counter(), {}
    for chunk in _chunks(jd):
        prev = None
        for t in _tokens(chunk):
            if t in STOPWORDS or len(t) < 2:
                prev = None
                continue
            s = stem(t)
            surface.setdefault(s, t)
            if (prev is not None and prev not in skip and s not in skip
                    and not (prev in GENERIC and s in GENERIC)):
                bi[(prev, s)] += 1
            if s not in GENERIC and s not in skip:
                uni[s] += 1
            prev = s
    out = [{"key": p, "label": f"{surface[p[0]]} {surface[p[1]]}", "count": c}
           for p, c in bi.most_common() if c >= 2][:n_bi]
    out += [{"key": s, "label": surface[s], "count": c} for s, c in uni.most_common(n_uni)]
    return out


def _vector(text):
    counts = Counter()
    for chunk in _chunks(text):
        prev = None
        for t in _tokens(chunk):
            s = stem(t)
            if t in STOPWORDS or s in GENERIC or len(t) < 2:
                prev = None
                continue
            counts[s] += 1
            if prev is not None:
                counts[(prev, s)] += 1
            prev = s
    return {k: 1 + math.log(c) for k, c in counts.items()}


def cosine(a, b):
    va, vb = _vector(a), _vector(b)
    dot = sum(w * vb.get(k, 0.0) for k, w in va.items())
    na = math.sqrt(sum(w * w for w in va.values()))
    nb = math.sqrt(sum(w * w for w in vb.values()))
    return dot / (na * nb) if na and nb else 0.0


# --------------------------------------------------------------- scoring ----
QUANT_RE = re.compile(r"\d+(?:\.\d+)?\s?%|\$\s?\d|\b\d+[kKmM]\+?(?![A-Za-z])")


def _rating(score):
    if score >= 85: return "Excellent match"
    if score >= 70: return "Strong match"
    if score >= 55: return "Good match"
    if score >= 40: return "Partial match"
    return "Weak match"


def analyze(jd, resume):
    sections = split_sections(resume)
    contact = extract_contact(resume)
    words = len(resume.split())

    # 1. skills (40)
    jd_sk, rs_sk = find_skills(jd), find_skills(resume)
    matched = [s for s in jd_sk if s in rs_sk]
    missing = [s for s in jd_sk if s not in rs_sk]
    extra = [s for s in rs_sk if s not in jd_sk]

    # 2. keywords (20), skipping words that are already part of a skill
    skip = {stem(t) for s in matched + missing for a in SKILLS[s] for t in _tokens(a)}
    keywords = extract_keywords(jd, skip)
    r_uni, r_bi = _stem_sets(resume)
    kw_hit, kw_miss, total_w, hit_w = [], [], 0, 0
    for k in keywords:
        hit = k["key"] in (r_bi if isinstance(k["key"], tuple) else r_uni)
        total_w += k["count"]
        if hit:
            hit_w += k["count"]
            kw_hit.append(k["label"])
        else:
            kw_miss.append(k["label"])
    kw_ratio = hit_w / total_w if total_w else 0.0

    if jd_sk:
        skill_ratio = len(matched) / len(jd_sk)
        skills_detail = f"{len(matched)} of {len(jd_sk)} skills from the job description found"
    else:
        skill_ratio = kw_ratio
        skills_detail = "No known skills found in the job description, so key terms were used instead"
    skills_pts = min(1.0, skill_ratio / 0.85) * 40   # matching 85% of a wish list counts as full marks
    kw_pts = min(1.0, kw_ratio / 0.70) * 20
    kw_detail = f"{len(kw_hit)} of {len(keywords)} key terms found" if keywords else "No key terms found in the job description"

    # 3. overall wording similarity (15)
    cos = cosine(jd, resume)
    sim_pts = min(1.0, cos / 0.40) * 15
    sim_detail = f"Wording similarity {cos * 100:.0f}%"

    # 4. experience (15)
    required = jd_required_years(jd)
    years = estimate_experience_years(resume, sections)
    if required:
        exp_pts = min(1.0, years / required) * 15
        exp_detail = f"About {years:g} years found; the job asks for {required}+"
    else:
        exp_pts = 15.0 if years > 0 else 9.0
        exp_detail = f"About {years:g} years found; the job does not state a minimum" if years else "No dated experience found"

    # 5. education (5)
    jd_deg = degree_levels(jd)
    req_level, req_label = (min(jd_deg) if jd_deg else (0, None))
    edu_text = sections.get("education") or resume
    res_deg = degree_levels(edu_text) or degree_levels(resume)
    res_level, res_label = (max(res_deg) if res_deg else (0, None))
    if req_level:
        if res_level >= req_level:
            edu_pts, edu_detail = 5.0, f"{res_label} found; the job asks for {req_label}"
        elif res_level:
            edu_pts, edu_detail = 2.5, f"{res_label} found; the job asks for {req_label}"
        else:
            edu_pts, edu_detail = 0.0, f"No degree found; the job asks for {req_label}"
    else:
        edu_pts = 5.0 if res_level else 3.0
        edu_detail = f"{res_label} found" if res_level else "No degree found"

    # 6. resume structure (5)
    checks = [
        bool(contact["email"]), bool(contact["phone"]),
        "experience" in sections, "education" in sections, "skills" in sections,
    ]
    struct_pts = float(sum(checks))
    struct_detail = f"{sum(checks)} of 5 basics present (email, phone, experience, education, skills)"

    breakdown = [
        {"key": "skills", "label": "Skills match", "score": skills_pts, "max": 40, "detail": skills_detail},
        {"key": "keywords", "label": "Key terms", "score": kw_pts, "max": 20, "detail": kw_detail},
        {"key": "similarity", "label": "Overall similarity", "score": sim_pts, "max": 15, "detail": sim_detail},
        {"key": "experience", "label": "Experience", "score": exp_pts, "max": 15, "detail": exp_detail},
        {"key": "education", "label": "Education", "score": edu_pts, "max": 5, "detail": edu_detail},
        {"key": "structure", "label": "Resume basics", "score": struct_pts, "max": 5, "detail": struct_detail},
    ]
    total = max(0.0, min(100.0, sum(b["score"] for b in breakdown)))
    for b in breakdown:
        b["score"] = round(b["score"], 1)

    tips = []
    if missing:
        tips.append("Add these skills if you have them, and show where you used them: " + ", ".join(missing[:8]) + ".")
    if kw_miss:
        tips.append("Use these terms from the job description where they are true: " + ", ".join(kw_miss[:6]) + ".")
    if required and years < required:
        tips.append(f"The job asks for about {required} years of experience and the resume shows about {years:g}. Give clear start and end dates for each role.")
    if req_level and res_level < req_level:
        tips.append(f"The job asks for a {req_label} degree. Add your education details if you have them.")
    for key, label in (("experience", "work experience"), ("education", "education"), ("skills", "skills")):
        if key not in sections:
            tips.append(f"Add a clearly labelled {label} section so screening tools can find it.")
    if not contact["email"]:
        tips.append("Add an email address.")
    if words < 250:
        tips.append(f"The resume is short ({words} words). Add detail about projects and results.")
    elif words > 1200:
        tips.append(f"The resume is long ({words} words). Trim it to what is relevant to this job.")
    if len(QUANT_RE.findall(resume)) < 3:
        tips.append("Add measurable results, such as 'cut page load time by 40%'.")

    parts = []
    if jd_sk:
        parts.append(f"The resume shows {len(matched)} of the {len(jd_sk)} skills the job asks for")
    parts.append(f"covers {len(kw_hit)} of {len(keywords)} key terms" if keywords else "has few key terms to compare")
    summary = ", and ".join(parts) + "."

    forms = {f for s in matched for f in rs_sk[s]} | set(kw_hit)
    return {
        "score": round(total, 1),
        "stars": round(total / 10) / 2,
        "rating": _rating(total),
        "summary": summary,
        "breakdown": breakdown,
        "skills": {"matched": matched, "missing": missing, "extra": extra[:12]},
        "keywords": {"matched": kw_hit, "missing": kw_miss},
        "experience": {"required_years": required, "detected_years": years},
        "education": {"required": req_label, "found": res_label},
        "candidate": contact,
        "suggestions": tips,
        "word_count": words,
        "highlight_terms": sorted(forms, key=len, reverse=True)[:300],
        "resume_text": resume[:8000],
        "resume_truncated": len(resume) > 8000,
    }


# ------------------------------------------------------------------ API ----
app = FastAPI(title="Resume scorer")


async def _read(upload):
    data = await upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "File is larger than 5 MB.")
    return data


@app.post("/api/analyze")
async def analyze_endpoint(
    resume: UploadFile = File(...),
    jd_text: str = Form(""),
    jd_file: Optional[UploadFile] = File(None),
):
    try:
        resume_text = extract_text(resume.filename, await _read(resume))
        if jd_file is not None and jd_file.filename:
            jd = extract_text(jd_file.filename, await _read(jd_file))
        else:
            jd = normalise(jd_text)
    except ParseError as exc:
        raise HTTPException(422, str(exc))
    if len(jd.split()) < 15:
        raise HTTPException(422, "The job description is too short. Paste the full text or upload a file.")
    return analyze(jd, resume_text)


@app.get("/api/health")
def health():
    return {"ok": True}


app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")


@app.get("/")
def index():
    return FileResponse(BASE / "static" / "index.html")